from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
import requests

from core import (DataError, Report, bulletin_numbers, collect, company_from_extract,
                  pkd_filters, yesterday)
from contacts import enrich_contacts
from demo import demo_report
from exports import pdf_bytes, xlsx_bytes

TARGET = date(2026, 10, 27)


@pytest.fixture
def extract():
    return {"odpis": {
        "naglowekA": {"numerKRS": "0000123456", "rejestr": "P",
                      "dataRejestracjiWKRS": "27.10.2026", "dataOstatniegoWpisu": "27.10.2026"},
        "dane": {
            "dzial1": {"danePodmiotu": {"nazwa": "TEST ODPADY SP. Z O.O.",
                                       "identyfikatory": {"nip": "1234567890"}},
                       "siedzibaIAdres": {"adres": {"miejscowosc": "Łódź"}},
                       "adresPocztyElektronicznej": "biuro@example.invalid"},
            "dzial3": {"przedmiotDzialalnosci": {
                "przedmiotPrzewazajacejDzialalnosci": [
                    {"kodDzial": "62", "kodKlasa": "01", "kodPodklasa": "Z", "opis": "IT"}],
                "przedmiotPozostalejDzialalnosci": [
                    {"kodDzial": "38", "kodKlasa": "11", "kodPodklasa": "Z", "opis": "Odpady"}]}}
        }}}


def test_modified_company_is_excluded_even_when_last_entry_is_target(extract):
    extract["odpis"]["naglowekA"]["dataRejestracjiWKRS"] = "26.10.2026"
    assert company_from_extract(extract, "0000123456", TARGET) is None


def test_registration_date_is_not_last_entry_date(extract):
    extract["odpis"]["naglowekA"]["dataOstatniegoWpisu"] = "28.10.2026"
    company = company_from_extract(extract, "0000123456", TARGET)
    assert company.registration_date == "2026-10-27"
    assert company.pkd == ["38.11.Z — Odpady"]
    assert company.emails == ["biuro@example.invalid"]


def test_filter_all_vs_main_and_full_code(extract):
    assert company_from_extract(extract, "0000123456", TARGET, main_only=True) is None
    assert company_from_extract(extract, "0000123456", TARGET, pkd_filters("38.11.Z"))
    assert company_from_extract(extract, "0000123456", TARGET, pkd_filters("39")) is None


def test_missing_registration_date_cannot_be_treated_as_new(extract):
    del extract["odpis"]["naglowekA"]["dataRejestracjiWKRS"]
    with pytest.raises(DataError):
        company_from_extract(extract, "0000123456", TARGET)


def test_extract_identity_must_match_request(extract):
    with pytest.raises(DataError):
        company_from_extract(extract, "0000999999", TARGET)


def test_bulletin_deduplicates_and_rejects_unknown_schema():
    assert bulletin_numbers(["0000123456", "123456"]) == ["0000123456"]
    with pytest.raises(DataError):
        bulletin_numbers({"unexpected": []})


def test_warsaw_midnight_changes_yesterday():
    # In Poland this instant is already 28 October.
    assert yesterday(datetime(2026, 10, 27, 23, 30, tzinfo=timezone.utc)) == TARGET


class Response:
    def __init__(self, payload=None, content=b"", fail=False):
        self.payload, self.content, self.fail = payload, content, fail
    def json(self):
        return self.payload
    def raise_for_status(self):
        if self.fail:
            raise requests.HTTPError("fixture HTTP error")


class Client:
    def __init__(self, responses):
        self.responses = iter(responses)
    def get(self, *args, **kwargs):
        return next(self.responses)


def test_partial_fetch_is_marked_incomplete(extract, monkeypatch):
    monkeypatch.setattr("core.time.sleep", lambda _: None)
    report = collect(TARGET, client=Client([
        Response(["123456", "123457"]), Response(extract), Response(fail=True)]))
    assert len(report.companies) == 1
    assert report.checked == 1 and report.candidates == 2
    assert not report.complete
    assert len(report.errors) == 1


def test_bulletin_failure_is_not_empty_success():
    with pytest.raises(DataError):
        collect(TARGET, client=Client([Response(fail=True)]))


def test_search_candidates_keep_sources_and_never_become_registry_contacts(extract):
    company = company_from_extract(extract, "0000123456", TARGET)
    content = b'''<rss><channel><item><title>TEST ODPADY SP. Z O.O.</title>
      <description>NIP 1234567890 kontakt: web@example.invalid tel: +48 501 502 503</description>
      <link>https://example.invalid/kontakt</link></item>
      <item><title>Inna firma</title><description>other@example.invalid</description>
      <link>https://example.invalid/inna</link></item></channel></rss>'''
    enrich_contacts(company, Client([Response(content=content)]))
    assert company.email_candidates == ["web@example.invalid"]
    assert company.emails == ["biuro@example.invalid"]
    assert company.phone_candidates == ["+48 501 502 503"]
    assert len(company.sources) == 1
    assert company.sources[0]["match"] == "NIP/KRS"


def test_search_html_error_is_not_no_contacts(extract):
    company = company_from_extract(extract, "0000123456", TARGET)
    with pytest.raises(DataError):
        enrich_contacts(company, Client([Response(content=b"<html>Blocked</html>")]))


def test_excel_preserves_krs_and_formula_like_names():
    report = demo_report(TARGET)
    report.companies[0].name = '=HYPERLINK("https://example.invalid")'
    book = load_workbook(BytesIO(xlsx_bytes(report)), data_only=False)
    assert book["Firmy"]["A2"].data_type == "s"
    assert book["Firmy"]["B2"].value == "0000000000"
    assert book["Informacje"]["B1"].value == "DEMO — dane fikcyjne"
    assert "Źródła kontaktów" in book.sheetnames


def test_pdf_generates_with_polish_text_and_long_sources():
    report = demo_report(TARGET)
    report.companies[0].name = "Żółć Łódź — " * 50
    report.companies[0].krs_source = "https://example.invalid/" + "a" * 1000
    content = pdf_bytes(report)
    assert content.startswith(b"%PDF-")
    assert len(content) > 10000


def test_app_selected_subclasses_and_both_downloads(extract, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from unittest.mock import Mock
    company = company_from_extract(extract, "0000123456", TARGET)
    report = Report(TARGET.isoformat(), "2026-10-28T09:00:00+01:00",
                    ["38.11.Z", "27.20.Z"], False, candidates=1, checked=1, companies=[company])
    fetch = Mock(return_value=report)
    save = Mock()
    monkeypatch.setattr("core.collect", fetch)
    monkeypatch.setattr("core.save_report", save)
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py").run()
    assert not app.exception
    assert len(app.sidebar.checkbox) == 1
    assert app.sidebar.checkbox[0].label == "Szukaj kontaktów w internecie"
    assert len(app.sidebar.button) == 1
    assert not app.sidebar.text_input
    labels = app.sidebar.multiselect[0].options
    assert "27.20.Z — Produkcja baterii i akumulatorów" in labels
    assert "38.21.Z — Odzysk surowców" in labels
    app.sidebar.date_input[0].set_value(TARGET)
    app.sidebar.multiselect[0].set_value(["38.11.Z", "27.20.Z"])
    app.sidebar.checkbox[0].uncheck()
    app.sidebar.button[0].click().run()
    assert not app.exception
    assert not app.session_state["report"].demo
    fetch.assert_called_once()
    assert fetch.call_args.args == (TARGET, ("38.11.Z", "27.20.Z"))
    assert fetch.call_args.kwargs["main_only"] is False
    assert fetch.call_args.kwargs["enrich"] is None
    save.assert_called_once()
    assert len(app.dataframe) == 1
    downloads = app.get("download_button")
    assert [button.proto.label for button in downloads] == [
        "Eksport do Excela (.xlsx)", "Eksport do PDF"]
    assert all(button.proto.url for button in downloads)


def test_app_empty_selection_does_not_fetch(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from unittest.mock import Mock
    fetch = Mock()
    monkeypatch.setattr("core.collect", fetch)
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py").run()
    app.sidebar.multiselect[0].set_value([]).run()
    assert not app.exception
    assert app.sidebar.button[0].disabled
    fetch.assert_not_called()


def test_demo_query_parameter_does_not_generate_fictional_results():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py")
    app.query_params["demo"] = "1"
    app.run()
    assert not app.exception
    assert not app.dataframe
    assert "report" not in app.session_state
